"""Tests for the L8-001 monthly ETF net-flow extractor."""

from datetime import date

import pytest
from openpyxl import Workbook

from goldrush2.dr2.extractors import l8_001


def observations(count=756, current=200.0, comparison=100.0):
    rows = [{"date": f"{2000 + i // 12:04d}-{i % 12 + 1:02d}-28", "value": 150.0} for i in range(count)]
    if count >= 5:
        rows[-5]["value"] = comparison
    rows[-1]["value"] = current
    return rows


@pytest.mark.parametrize("horizon, lookback", [("1-5d", 5), ("1-3m", 63), ("1-3y", 36), ("3-10y", 120)])
@pytest.mark.parametrize("current, comparison, signal", [(200.0, 100.0, 1), (50.0, 100.0, -1), (100.0, 100.0, 0)])
def test_directions_and_lookbacks(horizon, lookback, current, comparison, signal):
    rows = observations(current=current, comparison=100.0)
    rows[-lookback]["value"] = comparison
    result = l8_001.build_output(l8_001.VARIABLE_ID, l8_001.SOURCE_NAME, l8_001.SOURCE_URL, rows, as_of_date="2065-01-01", value_label="Gold ETF net flow")
    assert result["horizons"][horizon]["signal"] == signal
    assert result["horizons"][horizon]["confidence"] == 1


def test_flow_parser_sums_funds_and_excludes_total(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Demand by month"
    sheet.append(["Date", "Gold, US$/oz", "Ounces", "Tonnes", "Value (USD)", "Fund A", "Fund B"])
    sheet.append([date(2026, 7, 31), 3000, 0, 0, 9999, 10.5, -2.5])
    path = tmp_path / "ETF_Flows_test.xlsx"
    workbook.save(path)
    assert l8_001.parse_flows_workbook(path) == [{"date": "2026-07-31", "value": 8.0}]


def test_negative_flows_are_valid():
    result = l8_001.build_output(l8_001.VARIABLE_ID, l8_001.SOURCE_NAME, l8_001.SOURCE_URL, observations(current=-2, comparison=-1), as_of_date="2065-01-01", value_label="Gold ETF net flow")
    assert result["horizons"]["1-5d"]["signal"] == -1


def test_insufficient_history():
    result = l8_001.build_output(l8_001.VARIABLE_ID, l8_001.SOURCE_NAME, l8_001.SOURCE_URL, observations(4), as_of_date="2065-01-01", value_label="Gold ETF net flow")
    assert result["horizons"]["1-5d"]["confidence"] == 0


def test_missing_values_are_skipped(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Demand by month"
    sheet.append(["Date", "Value (USD)", "Fund A"])
    sheet.append(["bad", 0, 10])
    sheet.append([date(2026, 7, 31), 0, 10])
    path = tmp_path / "ETF_Flows_test.xlsx"
    workbook.save(path)
    assert len(l8_001.parse_flows_workbook(path)) == 1


def test_source_failure_output(monkeypatch, tmp_path):
    monkeypatch.setattr(l8_001.wgc, "fetch_wgc_workbook", lambda path: None)
    monkeypatch.setattr(l8_001.wgc, "LAST_FETCH_STALE", False)
    output = l8_001.run(output_path=tmp_path / "out.json", raw_dir=tmp_path)
    assert output["horizons"]["1-5d"]["confidence"] == 0
    assert "SOURCE UNAVAILABLE" in output["horizons"]["1-5d"]["evidence"]["summary"]


def test_stale_output(monkeypatch, tmp_path):
    monkeypatch.setattr(l8_001.wgc, "fetch_wgc_workbook", lambda path: None)
    monkeypatch.setattr(l8_001.wgc, "LAST_FETCH_STALE", True)
    output = l8_001.run(output_path=tmp_path / "out.json", raw_dir=tmp_path)
    assert "STALE DATA" in output["horizons"]["1-5d"]["evidence"]["summary"]


def test_schema():
    result = l8_001.build_output(l8_001.VARIABLE_ID, l8_001.SOURCE_NAME, l8_001.SOURCE_URL, observations(), as_of_date="2065-01-01", value_label="Gold ETF net flow")
    assert result["variable_id"] == "L8-001"
    assert result["data_frequency"] == "Monthly"
    assert set(result["horizons"]) == {"1-5d", "1-3m", "1-3y", "3-10y"}


def test_schedule_bound_t7_t8_crossing():
    rows = [{"date": f"2026-{month:02d}-28", "value": float(month)} for month in range(2, 7)]
    pending = l8_001.build_output(l8_001.VARIABLE_ID, l8_001.SOURCE_NAME, l8_001.SOURCE_URL, rows, as_of_date="2026-07-07")
    active = l8_001.build_output(l8_001.VARIABLE_ID, l8_001.SOURCE_NAME, l8_001.SOURCE_URL, rows, as_of_date="2026-07-08")
    assert pending["horizons"]["1-5d"]["status"] == "PENDING_RELEASE"
    assert pending["horizons"]["1-5d"]["confidence"] == 0
    assert active["horizons"]["1-5d"]["status"] == "VALID"
    assert active["horizons"]["1-5d"]["confidence"] == 1


def test_holdings_shaped_workbook_is_rejected(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Holdings"
    sheet.append(["Date", "Holdings (tonnes)"])
    sheet.append([date(2026, 6, 30), 100.0])
    path = tmp_path / "holdings.xlsx"
    workbook.save(path)
    with pytest.raises(ValueError, match="Demand by month"):
        l8_001.parse_flows_workbook(path)
